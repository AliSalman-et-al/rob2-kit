"""Researcher-facing Companion workspace projections derived from durable state."""

from __future__ import annotations

import json
from enum import StrEnum
from io import BytesIO
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from pydantic import Field, ValidationError

from rob2_kit.domain.assessment import (
    AlgorithmicJudgmentRevision,
    AssessmentRevision,
    JudgmentOverride,
    ReviewerProfileRevision,
)
from rob2_kit.domain.results import ResultSpecRevision
from rob2_kit.domain.revisions import FrozenModel, RecordReference
from rob2_kit.reports.archives import ArchiveVerificationError, verify_archive
from rob2_kit.storage.artifacts import ArtifactNotFoundError
from rob2_kit.storage.ledger import WorkflowEventOutcome, WorkflowLedger


class WorkspaceStage(StrEnum):
    READY = "ready"
    PREPARE = "prepare"
    REVIEW = "review"
    COMPLETE = "complete"


class SourcePreview(FrozenModel):
    source_id: str
    filename: str
    roles: tuple[str, ...]
    page_count: int
    custody_status: str
    parse_status: str
    preview_available: bool


class ResultCard(FrozenModel):
    result_id: str
    trial_id: str
    label: str
    identity: tuple[tuple[str, str], ...]
    sources: tuple[SourcePreview, ...] = ()
    column: str = "queue"
    checkpoint: str | None = None
    outcome: str | None = None
    reason: str | None = None
    next_action: str | None = None
    judgment: str | None = None


class WorkspaceProjection(FrozenModel):
    project_name: str
    stage: WorkspaceStage
    results: tuple[ResultCard, ...] = Field(min_length=1)
    selected_result_id: str
    agent_prompt: str
    connection_state: str
    review_locked: bool


class CompletionResultCard(FrozenModel):
    result_id: str | None = None
    trial_result: str
    identity: tuple[tuple[str, str], ...] = ()
    status: str
    assessment_revision_id: str | None = None
    sign_off_revision_id: str | None = None
    overall_judgment: str | None = None
    domain_judgments: tuple[tuple[str, str], ...] = ()
    reviewer: str | None = None
    reviewed_at: str | None = None
    limitation_count: int = 0
    artifact_readiness: str = "No signed Assessment"
    reason: str | None = None
    current: bool = False


class CompletionOutputCard(FrozenModel):
    title: str
    revision_binding: str
    state: str
    path: str | None = None
    limitation: str | None = None


class ProjectCompletion(FrozenModel):
    status: str
    results: tuple[CompletionResultCard, ...]
    outputs: tuple[CompletionOutputCard, ...]
    primary_action: str | None = None


def project_completion(
    ledger: WorkflowLedger,
    *,
    workspace: WorkspaceProjection | None,
    review_service: Any | None,
) -> ProjectCompletion | None:
    """Project progressive signed history and exact output bindings for Complete."""
    if review_service is None:
        return None
    sign_off = review_service.current_sign_off()
    materialization = review_service.output_materialization()
    events = ledger.events()
    sign_off_events = tuple(
        event
        for event in events
        if event.operation == "review:canonical-assessment-sign-off"
    )
    if not sign_off_events and not (
        workspace is not None
        and all(card.outcome is not None for card in workspace.results)
    ):
        return None
    withdrawal_ids = {
        str(_event_payload(ledger, event)["sign_off"]["revision_id"])
        for event in events
        if event.operation == "review:canonical-sign-off-withdrawal"
    }
    current_ids = {revision.revision_id for revision in ledger.current_revisions()}
    result_cards: list[CompletionResultCard] = []
    for event in sign_off_events:
        payload = _event_payload(ledger, event)
        is_current = (
            event.revision_id in current_ids
            and event.revision_id not in withdrawal_ids
        )
        status = (
            "Current signed Result"
            if is_current
            else "Withdrawn history"
            if event.revision_id in withdrawal_ids
            else "Invalidated history"
        )
        assessment = _assessment_details(ledger, payload["assessment"])
        event_materialization = _output_payload_for_sign_off(
            ledger, events, event.revision_id
        )
        result_cards.append(
            CompletionResultCard(
                result_id=assessment["result_id"],
                trial_result=assessment["trial_result"],
                identity=assessment["identity"],
                status=status,
                assessment_revision_id=str(payload["assessment"]["revision_id"]),
                sign_off_revision_id=event.revision_id,
                overall_judgment=assessment["overall_judgment"],
                domain_judgments=assessment["domain_judgments"],
                reviewer=_reviewer_display_name(ledger, payload),
                reviewed_at=str(payload["observed_at"]),
                limitation_count=assessment["limitation_count"],
                artifact_readiness=(
                    "Generated and current"
                    if is_current
                    and event_materialization is not None
                    and event_materialization.get("status") == "complete"
                    else "Signed — outputs incomplete"
                    if is_current
                    else "Immutable historical artifacts"
                    if event_materialization is not None
                    and event_materialization.get("status") == "complete"
                    else "Historical outputs incomplete"
                ),
                current=is_current,
            )
        )
    signed_result_ids = {
        card.result_id
        for card in result_cards
        if card.current and card.result_id is not None
    }
    if workspace is not None:
        for card in workspace.results:
            label = f"{card.trial_id} × {card.label}"
            if card.result_id in signed_result_ids:
                continue
            result_cards.append(
                CompletionResultCard(
                    result_id=card.result_id,
                    trial_result=label,
                    identity=card.identity,
                    status=_humanize(card.outcome or "not terminal"),
                    reason=card.reason,
                )
            )

    binding = (
        sign_off.revision_id if sign_off is not None else "No current Result sign-off"
    )
    paths = tuple(materialization.paths) if materialization is not None else ()
    output_specs = (
        ("Current report", (".html", ".md"), None),
        ("Structured export", (".json", ".csv"), None),
        ("Project workbook", (".xlsx",), None),
        ("Complete Verification archive", (".complete.rob2.zip",), None),
        (
            "Reference Verification archive",
            (".reference.rob2.zip",),
            "source integrity not independently verifiable",
        ),
    )
    outputs = tuple(
        _completion_output_card(
            title=title,
            suffixes=suffixes,
            limitation=limitation,
            binding=binding,
            paths=paths,
            assessment_revision_id=(
                review_service.review_case.assessment_revision_id
                if sign_off is not None
                else None
            ),
            generation_failed=bool(
                materialization is not None and materialization.failure_reason
            ),
        )
        for title, suffixes, limitation in output_specs
    )
    all_terminal = workspace is None or all(
        (
            card.result_id in signed_result_ids
            if card.outcome == "draft_ready"
            else card.outcome
            in {"excluded", "preparation_incomplete", "trial_failed", "failed"}
        )
        for card in workspace.results
    )
    no_deferred = not review_service.has_deferred_review()
    outputs_ready = bool(outputs) and all(
        output.state in {"Generated and current", "Verified and current"}
        for output in outputs
    )
    complete = all_terminal and no_deferred and sign_off is not None and outputs_ready
    status = (
        "Project Complete"
        if complete
        else "Signed — outputs incomplete"
        if sign_off is not None
        else "Completion in progress"
    )
    return ProjectCompletion(
        status=status,
        results=tuple(result_cards),
        outputs=outputs,
        primary_action=(
            materialization.retry_action
            if materialization is not None and materialization.failure_reason
            else None
        ),
    )


def _assessment_details(
    ledger: WorkflowLedger, reference: dict[str, Any]
) -> dict[str, Any]:
    """Resolve display facts only from one sign-off's immutable Assessment graph."""
    unavailable = {
        "result_id": None,
        "trial_result": str(reference.get("entity_id", "Assessment unavailable")),
        "identity": (),
        "overall_judgment": "Not available",
        "domain_judgments": tuple(
            (f"D{index}", "Not available") for index in range(1, 6)
        ),
        "limitation_count": 0,
    }
    try:
        assessment = AssessmentRevision.model_validate_json(
            ledger.artifacts.read(str(reference["content_hash"]))
        )
        result_spec = ResultSpecRevision.model_validate_json(
            ledger.artifacts.read(assessment.result_spec.content_hash)
        )
        judgments = tuple(
            AlgorithmicJudgmentRevision.model_validate_json(
                ledger.artifacts.read(item.content_hash)
            )
            for item in assessment.judgments
        )
        overrides = tuple(
            JudgmentOverride.model_validate_json(
                ledger.artifacts.read(item.content_hash)
            )
            for item in assessment.judgment_overrides
        )
    except (
        ArtifactNotFoundError,
        KeyError,
        ValidationError,
        ValueError,
    ):
        return unavailable
    replacements = {
        item.judgment_revision.revision_id: item.replacement.value
        for item in overrides
    }
    by_domain = {
        item.domain_id: replacements.get(item.revision_id, item.judgment.value)
        for item in judgments
    }
    displayed = tuple(
        (
            f"D{index}",
            _humanize(by_domain.get(f"domain:{index}", "not available")),
        )
        for index in range(1, 6)
    )
    rank = {"low": 0, "some_concerns": 1, "high": 2}
    overall = (
        _humanize(max(by_domain.values(), key=lambda item: rank.get(item, -1)))
        if by_domain
        else "Not available"
    )
    result = result_spec.result
    return {
        "result_id": result.result_id,
        "trial_result": (
            f"{result.trial_id} × {result.outcome_construct} · {result.time_point}"
        ),
        "identity": (
            ("Result", result.result_id),
            ("Trial", result.trial_id),
            ("Randomization", result.randomization_id),
            (
                "Comparison",
                f"{result.comparison.experimental_arm_id} vs "
                f"{result.comparison.comparator_arm_id}",
            ),
            ("Effect of interest", result.effect_of_interest),
            ("Outcome", result.outcome_construct),
            ("Measurement", result.measurement_instrument),
            ("Time point", result.time_point),
            ("Analysis population", result.analysis_population),
            ("Analysis model", result.analysis_model),
            ("Effect measure", result.effect_measure),
            ("Source locator", result.source_locator),
        ),
        "overall_judgment": overall,
        "domain_judgments": displayed,
        "limitation_count": len(assessment.review_findings),
    }


def _output_payload_for_sign_off(
    ledger: WorkflowLedger, events: tuple[Any, ...], sign_off_revision_id: str
) -> dict[str, Any] | None:
    for event in reversed(events):
        if event.operation != "review:materialize-signed-outputs":
            continue
        payload = _event_payload(ledger, event)
        if payload.get("sign_off", {}).get("revision_id") == sign_off_revision_id:
            return payload
    return None


def _reviewer_display_name(
    ledger: WorkflowLedger, sign_off: dict[str, Any]
) -> str:
    reference = sign_off.get("reviewer_profile", {})
    try:
        profile = ReviewerProfileRevision.model_validate_json(
            ledger.artifacts.read(str(reference["content_hash"]))
        )
    except (ArtifactNotFoundError, KeyError, ValidationError, ValueError):
        return "Attribution unavailable"
    return profile.display_name


def _completion_output_card(
    *,
    title: str,
    suffixes: tuple[str, ...],
    limitation: str | None,
    binding: str,
    paths: tuple[str, ...],
    assessment_revision_id: str | None,
    generation_failed: bool,
) -> CompletionOutputCard:
    path = next((item for item in paths if item.endswith(suffixes)), None)
    state = "Generation failed" if generation_failed else "Not generated"
    if path is not None and Path(path).is_file():
        if path.endswith((".complete.rob2.zip", ".reference.rob2.zip")):
            try:
                receipt = verify_archive(Path(path).read_bytes())
            except (ArchiveVerificationError, OSError):
                state = "Verification failed"
            else:
                state = (
                    "Verified and current"
                    if receipt.assessment_revision_id == assessment_revision_id
                    else "Revision binding failed"
                )
        else:
            state = (
                "Verified and current"
                if assessment_revision_id is not None
                and _derived_output_binds_assessment(
                    Path(path), assessment_revision_id
                )
                else "Revision binding failed"
            )
    return CompletionOutputCard(
        title=title,
        revision_binding=binding,
        state=state,
        path=path,
        limitation=limitation,
    )


def _derived_output_binds_assessment(
    path: Path, assessment_revision_id: str
) -> bool:
    try:
        if path.suffix.casefold() == ".xlsx":
            workbook = load_workbook(
                BytesIO(path.read_bytes()),
                read_only=True,
                data_only=True,
            )
            try:
                return any(
                    cell.value == assessment_revision_id
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                )
            finally:
                workbook.close()
        content = path.read_bytes()
        if path.suffix.casefold() == ".json":
            return json.loads(content).get("assessment_revision_id") == assessment_revision_id
        return assessment_revision_id.encode() in content
    except (OSError, UnicodeError, ValueError, json.JSONDecodeError):
        return False


class AuditEvidenceItem(FrozenModel):
    claim_id: str
    stance: str
    exact_text: str
    source_id: str
    source_artifact_hash: str
    page: int
    provenance: str
    spatial: tuple[float, float, float, float] | None = None
    source_conflict: bool = False
    visual_transcription: str | None = None
    visual_review_required: bool = False
    visual_current: bool = True


class SignalingQuestionAudit(FrozenModel):
    result_id: str
    domain_id: str
    sq_id: str
    guidance: str
    answer: str
    rationale: str
    matched_rules: tuple[str, ...]
    judgment: str
    coverage_state: str
    coverage_limitations: tuple[str, ...]
    no_information_basis: bool
    evidence: tuple[AuditEvidenceItem, ...]
    selected_evidence_id: str
    assessment_current: bool
    view: str = "crop"


def project_workspace(
    ledger: WorkflowLedger,
    *,
    connection_state: str,
    selected_result_id: str | None = None,
) -> WorkspaceProjection | None:
    """Join the current Companion view from ledger events and immutable artifacts."""
    events = ledger.events()
    initialization_event = next(
        (event for event in events if event.operation == "operation:initialize-project"),
        None,
    )
    if initialization_event is None:
        return None
    return _project_workspace_from_initialization(
        ledger,
        events,
        initialization_event,
        connection_state=connection_state,
        selected_result_id=selected_result_id,
    )


def signaling_question_audit(
    ledger: WorkflowLedger,
    *,
    result_id: str,
    domain_id: str | None,
    sq_id: str | None,
    evidence_id: str | None,
    assessment_current: bool,
    assessment: RecordReference | None,
    view: str = "crop",
) -> SignalingQuestionAudit | None:
    """Project one SQ audit from immutable preparation records and canonical units."""
    if assessment is None:
        return None
    events = ledger.events()
    try:
        assessment_payload = json.loads(ledger.artifacts.read(assessment.content_hash))
    except (ValueError, ArtifactNotFoundError):
        return None
    answer_references = tuple(assessment_payload.get("answers", ()))
    answer_events = tuple(
        event
        for reference in answer_references
        if (
            event := _referenced_event(
                events, reference, "operation:submit-sq-answers"
            )
        )
        is not None
    )
    result_key = result_id.split(":", 1)[-1]
    answer_event = next(
        (event for event in answer_events if event.scope == f"preparation:{result_key}"),
        None,
    )
    if answer_event is None:
        return None
    scope = answer_event.scope
    answer_payload = _event_payload(ledger, answer_event)
    answers = tuple(answer_payload.get("answers", ()))
    answer = next(
        (item for item in answers if sq_id is not None and item["question_id"] == sq_id),
        answers[0] if answers else None,
    )
    if answer is None:
        return None
    selected_sq = str(answer["question_id"])
    selected_domain = _domain_for_sq(selected_sq)
    if domain_id is not None and domain_id != selected_domain:
        return None

    bundle_event = next(
        (
            event
            for reference in assessment_payload.get("evidence_bundles", ())
            if (
                event := _referenced_event(
                    events, reference, "operation:freeze-evidence-bundle"
                )
            )
            is not None
            and event.scope == scope
        ),
        None,
    )
    if bundle_event is None:
        return None
    bundle = _event_payload(ledger, bundle_event)
    disposition_event = _referenced_event(
        events,
        bundle.get("disposition"),
        "operation:submit-evidence-dispositions",
    )
    judgment_event = _assessment_judgment(
        ledger,
        events,
        assessment_payload,
        answer_event,
        selected_domain,
    )
    trace_event = (
        _referenced_event(
            events,
            _event_payload(ledger, judgment_event).get("decision_trace"),
            "operation:derive-decision-trace",
        )
        if judgment_event is not None
        else None
    )
    trace = _event_payload(ledger, trace_event) if trace_event is not None else {}
    conflicts = {
        claim_id
        for conflict in bundle.get("conflicts", ())
        for claim_id in conflict
    }
    transcriptions = _visual_transcriptions(
        ledger, events, bundle_event, selected_sq
    )
    evidence = _audit_evidence(
        ledger,
        _event_payload(ledger, disposition_event) if disposition_event is not None else {},
        bundle,
        conflicts,
        transcriptions,
    )
    if not evidence:
        return None
    selected_evidence = (
        evidence_id
        if evidence_id in {item.claim_id for item in evidence}
        else evidence[0].claim_id
    )
    return SignalingQuestionAudit(
        result_id=result_id,
        domain_id=selected_domain,
        sq_id=selected_sq,
        guidance=_guidance_wording(selected_sq),
        answer=_humanize(str(answer["answer"])),
        rationale=str(answer["rationale"]),
        matched_rules=tuple(str(item) for item in trace.get("matched_rule_ids", ())),
        judgment=_humanize(str(trace.get("resulting_judgment", "not available"))),
        coverage_state=_humanize(str(bundle.get("coverage_state", "not recorded"))),
        coverage_limitations=tuple(
            str(item) for item in bundle.get("coverage_limitations", ())
        ),
        no_information_basis=bool(bundle.get("no_information_basis", False)),
        evidence=evidence,
        selected_evidence_id=selected_evidence,
        assessment_current=assessment_current,
        view="page" if view == "page" else "crop",
    )


def _audit_evidence(
    ledger: WorkflowLedger,
    dispositions_payload: dict[str, Any],
    bundle: dict[str, Any],
    conflicts: set[str],
    transcriptions: tuple[dict[str, Any], ...],
) -> tuple[AuditEvidenceItem, ...]:
    stances = {
        str(claim_id): (
            "Context"
            if disposition["kind"] == "accepted_contextual"
            else _humanize(str(disposition["kind"]).removeprefix("accepted_"))
        )
        for disposition in dispositions_payload.get("dispositions", ())
        if str(disposition.get("kind", "")).startswith("accepted_")
        for claim_id in disposition.get("claim_ids", ())
    }
    items: list[AuditEvidenceItem] = []
    for reference in bundle.get("items", ()):
        try:
            claim = json.loads(ledger.artifacts.read(str(reference["content_hash"])))
        except (KeyError, ValueError):
            continue
        claim_id = str(claim.get("claim_id", claim.get("entity_id", "")))
        if claim_id not in stances or not claim.get("quoted_text"):
            continue
        transcription = next(
            (
                item
                for item in transcriptions
                if item.get("source_id") == claim.get("source_id")
                and _visual_page(item) == int(claim.get("page", 0))
            ),
            None,
        )
        items.append(
            AuditEvidenceItem(
                claim_id=claim_id,
                stance=stances[claim_id],
                exact_text=str(claim["quoted_text"]),
                source_id=str(claim["source_id"]),
                source_artifact_hash=str(claim["source_artifact_hash"]),
                page=int(claim["page"]),
                provenance=(
                    f"{claim['parse_id']} · {claim['source_artifact_hash']} · "
                    f"span {claim['span_start']}:{claim['span_end']}"
                ),
                spatial=claim.get("spatial"),
                source_conflict=claim_id in conflicts,
                visual_transcription=(
                    str(transcription["transcription"]) if transcription else None
                ),
                visual_review_required=bool(
                    transcription and transcription.get("review_required", True)
                ),
                visual_current=bool(
                    transcription is None or transcription.get("_current", False)
                ),
            )
        )
    return tuple(items)


def _visual_transcriptions(
    ledger: WorkflowLedger,
    events: tuple[Any, ...],
    bundle_event: Any,
    sq_id: str,
) -> tuple[dict[str, Any], ...]:
    items = []
    current = {
        (revision.entity_id, revision.revision_id)
        for revision in ledger.current_revisions()
    }
    seen: set[tuple[str, int]] = set()
    bound = {
        (dependency.revision_id, dependency.content_hash)
        for dependency in bundle_event.dependencies
    }
    for event in reversed(events):
        if (
            event.operation != "operation:submit-visual-transcription"
            or (event.revision_id, event.output_revision_hashes[0]) not in bound
        ):
            continue
        transcription = _event_payload(ledger, event).get("transcription")
        if not transcription or (
            transcription.get("sq_id")
            and transcription["sq_id"] != sq_id
        ) or (
            not transcription.get("sq_id")
            and transcription.get("decision_critical")
            and sq_id not in transcription["decision_critical"]
        ):
            continue
        key = (str(transcription.get("source_id")), _visual_page(transcription))
        if key in seen:
            continue
        seen.add(key)
        items.append(
            {
                **transcription,
                "_current": (event.entity_id, event.revision_id) in current,
            }
        )
    return tuple(items)


def _visual_page(transcription: dict[str, Any]) -> int:
    render = transcription.get("render") or {}
    return int(transcription.get("page", render.get("page", 0)))


def _dependency_event(
    events: tuple[Any, ...], record_event: Any, operation: str
) -> Any | None:
    bound = {
        (dependency.revision_id, dependency.content_hash)
        for dependency in record_event.dependencies
    }
    return next(
        (
            event
            for event in reversed(events)
            if event.operation == operation
            and (event.revision_id, event.output_revision_hashes[0]) in bound
        ),
        None,
    )


def _assessment_judgment(
    ledger: WorkflowLedger,
    events: tuple[Any, ...],
    assessment: dict[str, Any],
    answer_event: Any,
    domain_id: str,
) -> Any | None:
    answer_hash = answer_event.output_revision_hashes[0]
    for reference in assessment.get("judgments", ()):
        event = _referenced_event(events, reference, "operation:derive-judgment")
        if event is None:
            continue
        payload = _event_payload(ledger, event)
        if payload.get("domain_id") != domain_id:
            continue
        if any(
            reference.get("revision_id") == answer_event.revision_id
            and reference.get("content_hash") == answer_hash
            for reference in payload.get("answer_revisions", ())
        ):
            return event
    return None


def _referenced_event(
    events: tuple[Any, ...], reference: dict[str, Any] | None, operation: str
) -> Any | None:
    if not reference:
        return None
    return next(
        (
            event
            for event in reversed(events)
            if event.operation == operation
            and event.revision_id == reference.get("revision_id")
            and event.output_revision_hashes[0] == reference.get("content_hash")
        ),
        None,
    )


def _event_payload(ledger: WorkflowLedger, event: Any) -> dict[str, Any]:
    return json.loads(ledger.artifacts.read(event.output_revision_hashes[0]))


def _domain_for_sq(sq_id: str) -> str:
    parts = sq_id.split(":")
    return f"domain:{parts[1]}" if len(parts) > 2 else "domain:unknown"


def _guidance_wording(sq_id: str) -> str:
    package = Path(__file__).resolve().parents[1]
    repository = Path(__file__).resolve().parents[3]
    path = next(
        candidate
        for candidate in (
            package / "packs" / "guidance" / "rob2-parallel-assignment-en-2019.1.yaml",
            repository / "packs" / "guidance" / "rob2-parallel-assignment-en-2019.1.yaml",
        )
        if candidate.is_file()
    )
    pack = yaml.safe_load(path.read_text(encoding="utf-8"))
    return next(
        (
            str(item["text"])
            for item in pack["items"]
            if item["logic_element_id"] == sq_id
        ),
        sq_id,
    )


def _project_workspace_from_initialization(
    ledger: WorkflowLedger,
    events: tuple[Any, ...],
    initialization_event: Any,
    *,
    connection_state: str,
    selected_result_id: str | None,
) -> WorkspaceProjection:
    payload = json.loads(ledger.artifacts.read(initialization_event.output_revision_hashes[0]))
    initialization = payload["initialization"]
    manifest = initialization["manifest"]
    plans = tuple(payload.get("preparation_plans", ()))
    trials = {trial["trial_id"]: trial for trial in initialization["trials"]}
    declared_specs = {
        item["result"]["result_id"]: item
        for item in initialization.get("result_specs", ())
    }
    planned_results = tuple(
        _result_card(
            ledger,
            trials[plan["trial_id"]],
            plan,
            events,
            declared_specs.get(plan.get("result_id")),
        )
        for plan in plans
    )
    planned_trial_ids = {plan["trial_id"] for plan in plans}
    unplanned_trials = tuple(
        trial
        for trial in initialization["trials"]
        if trial["trial_id"] not in planned_trial_ids
    )
    failed_results = tuple(
        _result_card(ledger, trial, None, events, result_spec)
        for trial in unplanned_trials
        for result_spec in (
            tuple(
                item
                for item in declared_specs.values()
                if item["result"]["trial_id"] == trial["trial_id"]
            )
            or (None,)
        )
    )
    results = (*planned_results, *failed_results)
    selected = (
        selected_result_id
        if selected_result_id in {item.result_id for item in results}
        else None
    )
    selected = selected or results[0].result_id
    all_terminal = all(item.outcome is not None for item in results)
    any_started = any(
        item.column == "processing"
        or item.outcome in {"draft_ready", "preparation_incomplete"}
        for item in results
    )
    any_reviewable = any(item.outcome == "draft_ready" for item in results)
    stage = (
        WorkspaceStage.REVIEW
        if all_terminal and any_reviewable
        else WorkspaceStage.PREPARE
        if any_started
        else WorkspaceStage.READY
    )
    root = manifest["input_root"]
    action = "Resume" if any_started else "Start"
    prompt = (
        f"{action} uninterrupted autonomous preparation for project "
        f"“{manifest['project_name']}” at {root}. Use continue_preparation until every "
        "requested Trial reaches a terminal Preparation outcome; do not ask mid-preparation "
        "questions."
    )
    return WorkspaceProjection(
        project_name=manifest["project_name"],
        stage=stage,
        results=results,
        selected_result_id=selected,
        agent_prompt=prompt,
        connection_state=connection_state,
        review_locked=not all_terminal,
    )


def _humanize(value: str) -> str:
    return value.replace("_", " ").capitalize()


def _result_card(
    ledger: WorkflowLedger,
    trial: dict[str, Any],
    plan: dict[str, Any] | None,
    events: tuple[Any, ...],
    declared_result_spec: dict[str, Any] | None = None,
) -> ResultCard:
    trial_id = str(trial["trial_id"])
    scope = (
        str(plan["scope"])
        if plan is not None
        else f"preparation:{trial_id.removeprefix('trial:')}"
    )
    plan_operations = {
        step["operation"] for step in (plan or {}).get("steps", ())
    }
    scope_events = tuple(
        event
        for event in events
        if event.scope == scope
        and (
            event.operation in plan_operations
            or event.operation == "operation:derive-assessment-record"
            or event.outcome
            in {
                WorkflowEventOutcome.PREPARATION_OUTCOME_REACHED,
                WorkflowEventOutcome.RETRYABLE_INTERRUPTION,
            }
        )
    )
    result_payload = declared_result_spec or _operation_payload(
        ledger,
        scope_events,
        "operation:submit-result-resolution",
    )
    result = result_payload.get("result", {}) if result_payload else {}
    result_id = str(result.get("result_id", f"result:{trial_id.removeprefix('trial:')}"))
    identity = _identity(result, trial_id)
    label = _result_label(result, trial_id)
    outcome, reason = _outcome(ledger, scope_events)
    if trial.get("status") == "trial_failed":
        failure = trial.get("failure") or {}
        outcome = "trial_failed"
        reason = str(failure.get("detail", "Required Trial material could not be prepared."))
    checkpoint = next(
        (event.checkpoint for event in reversed(scope_events) if event.checkpoint is not None),
        None,
    )
    column = "review_queue" if outcome is not None else "processing" if scope_events else "queue"
    return ResultCard(
        result_id=result_id,
        trial_id=trial_id,
        label=label,
        identity=identity,
        sources=tuple(
            _source_preview(source)
            for source in trial["inventory"]["sources"]
            if source.get("availability") == "acquired"
        ),
        column=column,
        checkpoint=checkpoint,
        outcome=outcome,
        reason=reason,
        next_action=_next_action(outcome, reason),
        judgment=(
            _committed_overall_judgment(ledger, scope_events)
            if outcome == "draft_ready"
            else None
        ),
    )


def _operation_payload(
    ledger: WorkflowLedger, events: tuple[Any, ...], operation: str
) -> dict[str, Any] | None:
    event = next((item for item in reversed(events) if item.operation == operation), None)
    return (
        json.loads(ledger.artifacts.read(event.output_revision_hashes[0]))
        if event is not None
        else None
    )


def _outcome(
    ledger: WorkflowLedger, events: tuple[Any, ...]
) -> tuple[str | None, str | None]:
    event = next(
        (
            item
            for item in reversed(events)
            if item.outcome is WorkflowEventOutcome.PREPARATION_OUTCOME_REACHED
        ),
        None,
    )
    if event is None:
        interruption = next(
            (
                item
                for item in reversed(events)
                if item.outcome is WorkflowEventOutcome.RETRYABLE_INTERRUPTION
            ),
            None,
        )
        reason = (
            "Preparation was interrupted safely; committed work is preserved."
            if interruption
            else None
        )
        return None, reason
    payload = json.loads(ledger.artifacts.read(event.output_revision_hashes[0]))
    return str(payload["outcome"]), payload.get("detail")


def _source_preview(source: dict[str, Any]) -> SourcePreview:
    coverage = source.get("coverage") or ()
    roles = tuple(str(role).replace("_", " ") for role in source.get("roles") or ())
    return SourcePreview(
        source_id=str(source["source_id"]),
        filename=str(source["title"]),
        roles=roles,
        page_count=len(coverage),
        custody_status=str(source["availability"]).replace("_", " "),
        parse_status=str(source["processing"]).replace("_", " "),
        preview_available=bool(
            source.get("artifact_hash")
            and str(source.get("title", "")).lower().endswith(".pdf")
        ),
    )


def secured_source_pdf(ledger: WorkflowLedger, source_id: str) -> bytes | None:
    """Return acquired PDF bytes for a source issued by this project."""
    initialization_event = next(
        (
            event
            for event in ledger.events()
            if event.operation == "operation:initialize-project"
        ),
        None,
    )
    if initialization_event is None:
        return None
    payload = json.loads(ledger.artifacts.read(initialization_event.output_revision_hashes[0]))
    for trial in payload["initialization"]["trials"]:
        for source in trial["inventory"]["sources"]:
            if (
                source["source_id"] == source_id
                and source.get("availability") == "acquired"
                and source.get("artifact_hash")
                and str(source.get("title", "")).lower().endswith(".pdf")
            ):
                return ledger.artifacts.read(source["artifact_hash"])
    return None


def _identity(result: dict[str, Any], trial_id: str) -> tuple[tuple[str, str], ...]:
    if not result:
        return (("Trial", trial_id), ("Result", "Pending exact Result resolution"))
    comparison = result.get("comparison") or {}
    return (
        ("Trial", trial_id),
        ("Randomization", str(result.get("randomization_id", ""))),
        ("Outcome", str(result.get("outcome_construct", ""))),
        ("Measurement", str(result.get("measurement_instrument", ""))),
        ("Time point", str(result.get("time_point", ""))),
        (
            "Comparison",
            f"{comparison.get('experimental_arm_id', '')} vs "
            f"{comparison.get('comparator_arm_id', '')}",
        ),
        ("Effect", str(result.get("effect_of_interest", ""))),
        ("Analysis population", str(result.get("analysis_population", ""))),
        ("Analysis model", str(result.get("analysis_model", ""))),
        ("Measure", str(result.get("effect_measure", ""))),
        ("Source locator", str(result.get("source_locator", ""))),
    )


def _result_label(result: dict[str, Any], trial_id: str) -> str:
    if not result:
        return f"{trial_id} · Result identity pending"
    return (
        f"{trial_id} · {result.get('outcome_construct', 'Result')} · "
        f"{result.get('time_point', 'time point pending')}"
    )


def _next_action(outcome: str | None, reason: str | None) -> str | None:
    if outcome == "preparation_incomplete":
        return "Resolve the material preparation gap, then resume preparation."
    if outcome == "trial_failed":
        return "Repair or replace the required source or Result information."
    if outcome is None and reason:
        return "Reconnect the agent and resume from the last durable checkpoint."
    return None


def _committed_overall_judgment(
    ledger: WorkflowLedger,
    events: tuple[Any, ...],
) -> str:
    assessment_event = next(
        (
            event
            for event in reversed(events)
            if event.operation == "operation:derive-assessment-record"
        ),
        None,
    )
    if assessment_event is None:
        return "Assessment committed"
    assessment = json.loads(
        ledger.artifacts.read(assessment_event.output_revision_hashes[0])
    )
    levels = tuple(
        json.loads(ledger.artifacts.read(reference["content_hash"]))["judgment"]
        for reference in assessment["judgments"]
    )
    overall = (
        "high"
        if "high" in levels
        else "some_concerns"
        if "some_concerns" in levels
        else "low"
    )
    return f"Overall judgment: {overall.replace('_', ' ').title()}"
